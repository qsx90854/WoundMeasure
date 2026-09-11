import csv
import json
from pathlib import Path
import tempfile
import unittest

import cv2
import numpy as np
from openpyxl import load_workbook

from Algorithm.block_height_accuracy import (
    BlockAccuracyConfig, BlockAccuracyReport, build_block_plan,
    measurement_record, summarize_block_records, block_mae_color,
)


class BlockAccuracyTests(unittest.TestCase):
    def test_mae_color_boundaries_and_missing_values(self):
        for value, expected in ((0, '#B7E4C7'), (0.4999, '#B7E4C7'),
                                (0.5, '#FFD6A5'), (1.0, '#FFD6A5'),
                                (1.0001, '#FFB3B3'), (12, '#FFB3B3'),
                                (None, None), (float('nan'), None),
                                (float('inf'), None), (-1, None)):
            with self.subTest(mae=value):
                self.assertEqual(block_mae_color(value), expected)

    def setUp(self):
        self.corners = [[100, 100], [700, 100], [700, 500], [100, 500]]
        # Use an explicit nine-point fixture; user tuning of the production
        # default (e.g. 5x5) must not change these expected test counts.
        self.plan = build_block_plan(self.corners, (600, 800), BlockAccuracyConfig(samples_per_axis=3))

    def record(self, index, height=None, reason=''):
        sample = self.plan['samples'][index]
        return measurement_record(sample, dict(height_display_mm=height,
            fail_reason=reason, u=sample['u'], v=sample['v'],
            height_reference_source='Shared Pattern Plane'), 2.5)

    def test_default_24_blocks_and_216_inset_points(self):
        self.assertEqual(len(self.plan['blocks']), 24)
        self.assertEqual(len(self.plan['samples']), 216)
        heights = [block['true_height_mm'] for block in self.plan['blocks']]
        np.testing.assert_allclose(heights, np.arange(12.5, 0.5, -0.5))
        first = self.plan['samples'][:9]
        np.testing.assert_allclose([point['u'] for point in first], [125, 150, 175] * 3)
        np.testing.assert_allclose([point['v'] for point in first], np.repeat([125, 150, 175], 3))
        self.assertEqual(self.plan['blocks'][6]['true_height_mm'], 9.5)
        self.assertEqual(self.plan['blocks'][-1]['true_height_mm'], 1.0)
        for block in self.plan['blocks']:
            selected = [p for p in self.plan['samples'] if p['block_id'] == block['block_id']]
            self.assertEqual(len(selected), 9)
            self.assertTrue(all(cv2.pointPolygonTest(np.float32(block['corners_px']),
                                                    (p['u'], p['v']), False) > 0 for p in selected))

    def test_projective_grid_roundtrips_to_model_coordinates(self):
        plan = build_block_plan([[150, 90], [700, 130], [600, 500], [80, 460]], (600, 800))
        pixels = np.array([[p['u'], p['v']] for p in plan['samples']], np.float64)
        model = cv2.perspectiveTransform(pixels[None], np.linalg.inv(np.array(plan['homography'])))[0]
        np.testing.assert_allclose(model, [[p['model_x_mm'], p['model_y_mm']] for p in plan['samples']], atol=1e-8)

    def test_invalid_corner_order_degeneracy_and_bounds_rejected(self):
        for corners in ([[100, 100], [700, 500], [700, 100], [100, 500]],
                        [[100, 100], [100, 500], [700, 500], [700, 100]],
                        [[100, 100]] * 4,
                        [[-1, 100], [700, 100], [700, 500], [100, 500]],
                        [[float('nan'), 100], [700, 100], [700, 500], [100, 500]]):
            with self.subTest(corners=corners), self.assertRaises(ValueError):
                build_block_plan(corners, (600, 800))

    def test_grid_size_inset_and_height_sequence_are_tunable(self):
        config = BlockAccuracyConfig(columns=2, rows=1, samples_per_axis=1,
                                     first_height_mm=5, height_step_mm=-2)
        plan = build_block_plan(self.corners, (600, 800), config)
        self.assertEqual(len(plan['samples']), 2)
        self.assertEqual([p['true_height_mm'] for p in plan['samples']], [5, 3])
        np.testing.assert_allclose([[p['u'], p['v']] for p in plan['samples']], [[250, 300], [550, 300]])
        with self.assertRaises(ValueError):
            build_block_plan(self.corners, (600, 800), BlockAccuracyConfig(inset_fraction=0.5))

    def test_signed_height_and_error_preserved_without_abs_or_rounding(self):
        record = self.record(0, -1.23456)
        self.assertEqual(record['wound_height_mm'], -1.23456)
        self.assertAlmostEqual(record['error_mm'], -13.73456)
        self.assertAlmostEqual(record['abs_error_mm'], 13.73456)

    def test_failed_points_are_excluded_not_replaced_with_zero(self):
        records = [self.record(0, 12), self.record(1, 14), self.record(2, None),
                   self.record(3, None, 'triangulation rejected')]
        blocks, summary = summarize_block_records(self.plan, records)
        self.assertEqual(blocks[0]['valid_count'], 2)
        self.assertEqual(blocks[0]['failed_count'], 2)
        self.assertEqual(blocks[0]['unmeasured_count'], 5)
        self.assertEqual(blocks[0]['mean_height_mm'], 13)
        self.assertEqual(blocks[0]['std_height_mm'], 1)
        self.assertEqual(blocks[0]['mean_height_error_mm'], 0.5)
        self.assertEqual(blocks[0]['mae_mm'], 1)
        self.assertAlmostEqual(blocks[0]['rmse_mm'], np.sqrt(1.25))
        self.assertEqual(summary['block_mean_mae_mm'], 0.5)
        self.assertIsNone(blocks[1]['mean_height_mm'])

    def test_nonfinite_or_missing_heights_are_failures(self):
        for height in (None, float('nan'), float('inf')):
            record = self.record(0, height)
            self.assertEqual(record['status'], 'failed')
            self.assertIsNone(record['error_mm'])
            self.assertTrue(record['failure_reason'])

    def test_valid_fused_readout_keeps_best_pair_failure_as_warning(self):
        record = self.record(0, 12.6, 'BEST pair rejected; EXTRA pair fused')
        self.assertEqual(record['status'], 'valid')
        self.assertEqual(record['failure_reason'], '')
        self.assertTrue(record['measurement_warning'])
        self.assertAlmostEqual(record['error_mm'], 0.1)

    def test_full_known_bias_report(self):
        records = [self.record(i, point['true_height_mm'] + 0.2)
                   for i, point in enumerate(self.plan['samples'])]
        blocks, summary = summarize_block_records(self.plan, records)
        self.assertEqual(summary['complete_block_count'], 24)
        self.assertEqual(summary['valid_count'], 216)
        self.assertAlmostEqual(summary['mae_mm'], 0.2)
        self.assertAlmostEqual(summary['block_mean_mae_mm'], 0.2)
        self.assertTrue(all(abs(row['mean_height_error_mm'] - 0.2) < 1e-9 for row in blocks))

    def test_incremental_csv_cancel_summary_and_unique_directories(self):
        with tempfile.TemporaryDirectory() as directory:
            report = BlockAccuracyReport(directory, self.plan, {'normal': np.array([0, 0, 1])})
            report.append(self.record(0, 12.6))
            # The point is readable from disk even before finishing/cancelling.
            with (report.directory / 'points.csv').open(encoding='utf-8-sig') as stream:
                saved = list(csv.DictReader(stream))
            self.assertEqual(len(saved), 1)
            self.assertAlmostEqual(float(saved[0]['error_mm']), 0.1)
            blocks, summary = report.finish('cancelled', 'test cancel')
            self.assertEqual(summary['unmeasured_count'], 215)
            data = json.loads((report.directory / 'summary.json').read_text(encoding='utf-8'))
            self.assertEqual(data['status'], 'cancelled')
            self.assertEqual(len(data['blocks']), 24)
            self.assertIsNone(data['blocks'][1]['mean_height_mm'])
            other = BlockAccuracyReport(directory, self.plan, {})
            self.assertNotEqual(report.directory, other.directory)
            other.finish('cancelled')
            with self.assertRaises(RuntimeError):
                report.append(self.record(1, 12))

    def test_out_of_order_records_rejected_and_partial_not_completed(self):
        with tempfile.TemporaryDirectory() as directory:
            report = BlockAccuracyReport(directory, self.plan, {})
            with self.assertRaises(ValueError):
                report.append(self.record(1, 12))
            report.finish()
            data = json.loads((report.directory / 'summary.json').read_text(encoding='utf-8'))
            self.assertEqual(data['status'], 'partial')

    def test_csv_floats_fit_excel_precision_and_json_keeps_full_values(self):
        with tempfile.TemporaryDirectory() as directory:
            report = BlockAccuracyReport(directory, self.plan, {})
            record = self.record(0, 12.282020401027337)
            report.append(record)
            blocks, _ = report.finish('cancelled')
            for filename, original_rows in (('points.csv', [record]), ('blocks.csv', blocks)):
                with (report.directory / filename).open(encoding='utf-8-sig') as stream:
                    saved = list(csv.DictReader(stream))
                self.assertEqual(len(saved), len(original_rows))
                for row, original in zip(saved, original_rows):
                    for key, value in original.items():
                        if isinstance(value, (float, np.floating)):
                            self.assertEqual(row[key], format(value, '.15g'))
                            mantissa = row[key].lower().split('e')[0]
                            digits = mantissa.lstrip('-+').replace('.', '').lstrip('0')
                            self.assertLessEqual(len(digits), 15)
                            self.assertAlmostEqual(float(row[key]), value, places=10)
                        elif value is None:
                            self.assertEqual(row[key], '')
                self.assertNotIn("'", (report.directory / filename).read_text(encoding='utf-8-sig'))
            self.assertLess(float(saved[0]['mean_error_mm']), 0)
            self.assertEqual(report.records[0]['wound_height_mm'], record['wound_height_mm'])
            data = json.loads((report.directory / 'summary.json').read_text(encoding='utf-8'))
            self.assertEqual(data['blocks'][0]['mean_height_mm'], record['wound_height_mm'])

    def test_xlsx_has_gt_mean_layout_numeric_data_and_two_charts(self):
        with tempfile.TemporaryDirectory() as directory:
            report = BlockAccuracyReport(directory, self.plan, {})
            for index, sample in enumerate(self.plan['samples']):
                report.append(self.record(index, sample['true_height_mm'] + 0.2))
            blocks, _ = report.finish()
            workbook_path = report.directory / 'blocks.xlsx'
            self.assertTrue(workbook_path.exists())
            workbook = load_workbook(workbook_path, data_only=False)
            self.addCleanup(workbook.close)
            self.assertEqual(workbook.sheetnames, ['Summary', 'Chart Data'])
            summary = workbook['Summary']
            self.assertEqual([summary.cell(1, column).value for column in range(1, 13)],
                             ['GT', 'Mean_H'] * 6)
            self.assertEqual(summary['A2'].value, 12.5)
            self.assertAlmostEqual(summary['B2'].value, blocks[0]['mean_height_mm'])
            self.assertEqual(summary['K5'].value, 1.0)
            self.assertAlmostEqual(summary['L5'].value, blocks[-1]['mean_height_mm'])
            self.assertIsInstance(summary['B2'].value, float)
            self.assertEqual(len(summary._charts), 2)
            self.assertEqual(summary._charts[0].title.tx.rich.p[0].r[0].t,
                             'MAE vs True Height')
            self.assertEqual(summary._charts[1].title.tx.rich.p[0].r[0].t,
                             'STD vs True Height')
            data = workbook['Chart Data']
            self.assertEqual(data['B2'].value, 1.0)
            self.assertEqual(data['B25'].value, 12.5)
            self.assertAlmostEqual(data['D2'].value, 0.2)
            self.assertAlmostEqual(data['E2'].value, 0.0)
            self.assertIsInstance(data['D2'].value, float)


if __name__ == '__main__':
    unittest.main()
