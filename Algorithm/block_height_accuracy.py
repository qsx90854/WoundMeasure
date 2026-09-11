"""Four-corner sampling and incremental, UI-free stepped-block accuracy reports.

The homography defines an approximate planar layout, NOT a reconstruction of
the raised block tops. The UI must preview the grid before measurements start.
No outlier removal or absolute-height conversion is applied to measured data.
"""
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
import csv
import json
import math
import uuid

import cv2
import numpy as np


def block_mae_color(mae_mm):
    """Display-only MAE bands in mm; missing/invalid values stay unfilled."""
    value = _finite_number(mae_mm)
    if value is None or value < 0:
        return None
    if value < 0.5:
        return '#B7E4C7'  # light green
    if value <= 1.0:
        return '#FFD6A5'  # light orange (includes both boundaries)
    return '#FFB3B3'      # light red


@dataclass(frozen=True)
class BlockAccuracyConfig:
    columns: int = 6
    rows: int = 4
    cell_width_mm: float = 20.0
    cell_height_mm: float = 20.0
    samples_per_axis: int = 5
    inset_fraction: float = 0.25
    first_height_mm: float = 12.5
    height_step_mm: float = -0.5

    def validate(self):
        for value in (self.columns, self.rows, self.samples_per_axis):
            if not isinstance(value, int) or value < 1:
                raise ValueError('Grid/sample counts must be positive integers')
        values = (self.cell_width_mm, self.cell_height_mm, self.inset_fraction,
                  self.first_height_mm, self.height_step_mm)
        if not np.all(np.isfinite(values)):
            raise ValueError('Block dimensions/heights must be finite')
        if min(self.cell_width_mm, self.cell_height_mm) <= 0:
            raise ValueError('Block dimensions must be positive')
        if not 0 < self.inset_fraction < 0.5:
            raise ValueError('inset_fraction must lie strictly between 0 and 0.5')


def build_block_plan(corners, image_shape, config=BlockAccuracyConfig()):
    """Corners are TL, TR, BR, BL. Heights are row-major, left to right."""
    config.validate()
    quad = np.asarray(corners, dtype=np.float64)
    if quad.shape != (4, 2) or not np.all(np.isfinite(quad)):
        raise ValueError('Select four finite corners: TL, TR, BR, BL')
    height, width = image_shape[:2]
    if (np.any(quad < 0) or np.any(quad[:, 0] > width - 1)
            or np.any(quad[:, 1] > height - 1)):
        raise ValueError('All four corners must be inside the left image')
    edges = np.roll(quad, -1, axis=0) - quad
    turns = edges[:, 0] * np.roll(edges[:, 1], -1) - edges[:, 1] * np.roll(edges[:, 0], -1)
    if np.any(turns <= 1e-6) or cv2.contourArea(quad.astype(np.float32)) < 16:
        raise ValueError('Corners must form a non-crossing convex TL/TR/BR/BL quadrilateral')
    model_width = config.columns * config.cell_width_mm
    model_height = config.rows * config.cell_height_mm
    matrix = cv2.getPerspectiveTransform(
        np.array([[0, 0], [model_width, 0], [model_width, model_height],
                  [0, model_height]], np.float32), quad.astype(np.float32))

    def project(points):
        output = cv2.perspectiveTransform(np.asarray(points, np.float64).reshape(1, -1, 2), matrix)[0]
        if not np.all(np.isfinite(output)):
            raise ValueError('Degenerate corner projection')
        return output.tolist()

    fractions = (np.array([0.5]) if config.samples_per_axis == 1 else
                 np.linspace(config.inset_fraction, 1 - config.inset_fraction, config.samples_per_axis))
    samples, blocks = [], []
    for row in range(config.rows):
        for col in range(config.columns):
            block_id = row * config.columns + col + 1
            truth = config.first_height_mm + (block_id - 1) * config.height_step_mm
            x0, y0 = col * config.cell_width_mm, row * config.cell_height_mm
            corners_mm = [[x0, y0], [x0 + config.cell_width_mm, y0],
                          [x0 + config.cell_width_mm, y0 + config.cell_height_mm],
                          [x0, y0 + config.cell_height_mm]]
            blocks.append(dict(block_id=block_id, row=row + 1, column=col + 1,
                               true_height_mm=truth, corners_px=project(corners_mm),
                               center_px=project([[x0 + config.cell_width_mm / 2,
                                                   y0 + config.cell_height_mm / 2]])[0]))
            for sy, fy in enumerate(fractions):
                for sx, fx in enumerate(fractions):
                    x, y = x0 + fx * config.cell_width_mm, y0 + fy * config.cell_height_mm
                    u, v = project([[x, y]])[0]
                    samples.append(dict(sample_id=len(samples) + 1, block_id=block_id,
                                        row=row + 1, column=col + 1, sample_row=sy + 1,
                                        sample_column=sx + 1, model_x_mm=x, model_y_mm=y,
                                        u=u, v=v, true_height_mm=truth))
    return dict(config=asdict(config), corners_px=quad.tolist(), homography=matrix.tolist(),
                blocks=blocks, samples=samples,
                warning='Four-corner homography is approximate for non-coplanar stepped tops; inspect preview')


def _finite_number(value):
    try:
        return float(value) if value is not None and np.isfinite(float(value)) else None
    except (TypeError, ValueError):
        return None


def measurement_record(sample, result, elapsed_seconds):
    """Use the existing click pipeline's unrounded, signed Wound Height."""
    result = result or {}
    height = _finite_number(result.get('height_display_mm'))
    reason = result.get('fail_reason') or ''
    # height_display_mm is produced only by the accepted 3D/plane readout.
    # A failed BEST pair can leave fail_reason behind even when EXTRA pairs
    # supply a valid fused readout. Preserve that warning, not a false failure.
    valid = height is not None
    if not valid and not reason:
        reason = 'No valid Wound Height/reference plane'
    debug = result.get('region_debug') or {}
    record = dict(sample)
    record.update(measured_u=_finite_number(result.get('u')),
                  measured_v=_finite_number(result.get('v')),
                  wound_height_mm=height, error_mm=(height - sample['true_height_mm']) if valid else None,
                  abs_error_mm=abs(height - sample['true_height_mm']) if valid else None,
                  status='valid' if valid else 'failed',
                  failure_reason=str(reason) if not valid else '',
                  measurement_warning=str(reason) if valid else '',
                  height_reference_source=result.get('height_reference_source'),
                  method=result.get('method'), candidate_frame=result.get('cand_idx'),
                  fused_count=result.get('fused_count'),
                  group_score=_finite_number(debug.get('group_score')),
                  objective_score=_finite_number(debug.get('objective_score')),
                  elapsed_seconds=float(elapsed_seconds))
    return record


def _metrics(records):
    valid = [record for record in records if record['status'] == 'valid']
    heights = np.array([record['wound_height_mm'] for record in valid], dtype=float)
    errors = np.array([record['error_mm'] for record in valid], dtype=float)
    return dict(valid_count=len(valid), failed_count=len(records) - len(valid),
                mean_height_mm=float(np.mean(heights)) if len(valid) else None,
                std_height_mm=float(np.std(heights)) if len(valid) else None,
                mean_error_mm=float(np.mean(errors)) if len(valid) else None,
                mae_mm=float(np.mean(np.abs(errors))) if len(valid) else None,
                rmse_mm=float(np.sqrt(np.mean(errors ** 2))) if len(valid) else None)


def summarize_block_records(plan, records):
    blocks = []
    expected = plan['config']['samples_per_axis'] ** 2
    for block in plan['blocks']:
        selected = [record for record in records if record['block_id'] == block['block_id']]
        row = {key: block[key] for key in ('block_id', 'row', 'column', 'true_height_mm')}
        row.update(_metrics(selected))
        row.update(expected_count=expected, processed_count=len(selected),
                   unmeasured_count=expected - len(selected))
        row['mean_height_error_mm'] = row['mean_error_mm']
        row['abs_mean_height_error_mm'] = (abs(row['mean_error_mm'])
                                           if row['mean_error_mm'] is not None else None)
        blocks.append(row)
    summary = _metrics(records)
    summary.update(expected_count=len(plan['samples']), processed_count=len(records),
                   unmeasured_count=len(plan['samples']) - len(records),
                   valid_block_count=sum(row['valid_count'] > 0 for row in blocks),
                   complete_block_count=sum(row['valid_count'] == expected for row in blocks))
    block_errors = [row['mean_height_error_mm'] for row in blocks if row['valid_count']]
    summary['block_mean_mae_mm'] = float(np.mean(np.abs(block_errors))) if block_errors else None
    summary['block_mean_rmse_mm'] = float(np.sqrt(np.mean(np.square(block_errors)))) if block_errors else None
    return blocks, summary


def _json_safe(value):
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, np.ndarray)):
        return [_json_safe(item) for item in value]
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        return _finite_number(value)
    if isinstance(value, np.bool_):
        return bool(value)
    if isinstance(value, Path):
        return str(value)
    return value


def _csv_record(record):
    """Limit exported floats to Excel's 15 significant digits.

    Longer float reprs can be imported as text when Excel's automatic
    long-number conversion is disabled. Keep full precision in memory/JSON.
    """
    return {key: format(value, '.15g') if isinstance(value, (float, np.floating))
            else value for key, value in record.items()}


def write_blocks_xlsx(output_path, plan, blocks):
    """Write the 6x4 GT/mean view and two editable XY charts."""
    from openpyxl import Workbook
    from openpyxl.chart import Reference, ScatterChart, Series
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    output_path = Path(output_path)
    rows = int(plan['config']['rows'])
    columns = int(plan['config']['columns'])
    by_position = {(int(row['row']), int(row['column'])): row for row in blocks}

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = 'Summary'
    summary_sheet.sheet_view.showGridLines = False

    header_fill = PatternFill('solid', fgColor='4472C4')
    gt_fill = PatternFill('solid', fgColor='E7EAF2')
    mean_fill = PatternFill('solid', fgColor='D9E2F3')
    white_side = Side(style='thin', color='FFFFFF')
    cell_border = Border(left=white_side, right=white_side,
                         top=white_side, bottom=white_side)
    for column in range(1, columns + 1):
        gt_column = 2 * column - 1
        mean_column = 2 * column
        for target_column, label in ((gt_column, 'GT'), (mean_column, 'Mean_H')):
            cell = summary_sheet.cell(row=1, column=target_column, value=label)
            cell.fill = header_fill
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='left', vertical='center')
            cell.border = cell_border
            summary_sheet.column_dimensions[cell.column_letter].width = 10
        for row in range(1, rows + 1):
            block = by_position.get((row, column))
            gt_cell = summary_sheet.cell(row=row + 1, column=gt_column)
            mean_cell = summary_sheet.cell(row=row + 1, column=mean_column)
            if block is not None:
                gt_cell.value = _finite_number(block.get('true_height_mm'))
                mean_cell.value = _finite_number(block.get('mean_height_mm'))
            gt_cell.fill = gt_fill
            mean_cell.fill = mean_fill
            gt_cell.font = Font(name='Arial', size=10, bold=True, color='000000')
            mean_cell.font = Font(name='Arial', size=10, color='2F75B5')
            gt_cell.number_format = '0.##'
            mean_cell.number_format = '0.00'
            for cell in (gt_cell, mean_cell):
                cell.alignment = Alignment(horizontal='left', vertical='center')
                cell.border = cell_border
    summary_sheet.row_dimensions[1].height = 21
    for row in range(2, rows + 2):
        summary_sheet.row_dimensions[row].height = 24

    data_sheet = workbook.create_sheet('Chart Data')
    data_sheet.sheet_view.showGridLines = False
    data_headers = ('Block', 'True Height (mm)', 'Mean Height (mm)',
                    'MAE (mm)', 'STD (mm)', 'Valid Count')
    for column, label in enumerate(data_headers, 1):
        cell = data_sheet.cell(row=1, column=column, value=label)
        cell.fill = header_fill
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = cell_border
    sorted_blocks = sorted(blocks, key=lambda row: float(row['true_height_mm']))
    for row_index, block in enumerate(sorted_blocks, 2):
        values = (int(block['block_id']), _finite_number(block.get('true_height_mm')),
                  _finite_number(block.get('mean_height_mm')),
                  _finite_number(block.get('mae_mm')),
                  _finite_number(block.get('std_height_mm')),
                  int(block.get('valid_count', 0)))
        for column, value in enumerate(values, 1):
            cell = data_sheet.cell(row=row_index, column=column, value=value)
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(horizontal='right', vertical='center')
        for column in range(2, 6):
            data_sheet.cell(row=row_index, column=column).number_format = '0.000'
    for letter, width in {'A': 10, 'B': 19, 'C': 19, 'D': 13, 'E': 13, 'F': 14}.items():
        data_sheet.column_dimensions[letter].width = width
    data_sheet.freeze_panes = 'A2'
    data_sheet.auto_filter.ref = f'A1:F{len(sorted_blocks) + 1}'

    def add_scatter_chart(y_column, series_title, chart_title, y_axis_title,
                          anchor, color):
        chart = ScatterChart()
        chart.title = chart_title
        chart.style = 13
        chart.height = 7.2
        chart.width = 11.2
        chart.legend = None
        chart.x_axis.title = 'True Height (mm)'
        chart.y_axis.title = y_axis_title
        chart.x_axis.scaling.min = 0
        if sorted_blocks:
            max_height = max(float(row['true_height_mm']) for row in sorted_blocks)
            chart.x_axis.scaling.max = math.ceil(max_height + 1.0)
        chart.x_axis.majorUnit = 2
        chart.y_axis.scaling.min = 0
        chart.x_axis.numFmt = '0.0'
        chart.y_axis.numFmt = '0.00'
        max_row = len(sorted_blocks) + 1
        x_values = Reference(data_sheet, min_col=2, min_row=2, max_row=max_row)
        y_values = Reference(data_sheet, min_col=y_column, min_row=2, max_row=max_row)
        series = Series(y_values, x_values, title=series_title)
        series.marker.symbol = 'circle'
        series.marker.size = 5
        series.marker.graphicalProperties.solidFill = color
        series.marker.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 19050
        series.smooth = False
        chart.series.append(series)
        summary_sheet.add_chart(chart, anchor)

    add_scatter_chart(4, 'MAE', 'MAE vs True Height', 'MAE (mm)', 'A8', 'ED7D31')
    add_scatter_chart(5, 'STD', 'STD vs True Height', 'STD (mm)', 'G8', '4472C4')

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    temporary_path = output_path.with_name(
        f'.{output_path.stem}.{uuid.uuid4().hex[:8]}.tmp.xlsx')
    try:
        workbook.save(temporary_path)
        temporary_path.replace(output_path)
    finally:
        workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()


class BlockAccuracyReport:
    """One unique run directory. Each completed point is flushed immediately."""
    def __init__(self, output_root, plan, metadata):
        self.plan, self.records = plan, []
        self.directory = Path(output_root) / (datetime.now().strftime('%Y%m%d_%H%M%S') + '_' + uuid.uuid4().hex[:8])
        self.directory.mkdir(parents=True, exist_ok=False)
        self.metadata = _json_safe(metadata)
        self._stream = None
        self._writer = None
        self.closed = False
        self._write_json('run.json', dict(status='running', plan=plan, metadata=self.metadata))

    def _write_json(self, name, value):
        with (self.directory / name).open('w', encoding='utf-8') as stream:
            json.dump(_json_safe(value), stream, ensure_ascii=False, indent=2, allow_nan=False)

    def append(self, record):
        if self.closed:
            raise RuntimeError('Report is already closed')
        index = len(self.records)
        if index >= len(self.plan['samples']) or record['sample_id'] != self.plan['samples'][index]['sample_id']:
            raise ValueError('Measurement records must follow the planned sample order')
        if self._stream is None:
            self._stream = (self.directory / 'points.csv').open('x', newline='', encoding='utf-8-sig')
            self._writer = csv.DictWriter(self._stream, fieldnames=list(record))
            self._writer.writeheader()
        self._writer.writerow(_csv_record(record))
        self._stream.flush()
        self.records.append(dict(record))

    def finish(self, status='completed', reason=''):
        if self._stream is not None:
            self._stream.close()
        self.closed = True
        blocks, summary = summarize_block_records(self.plan, self.records)
        if status == 'completed' and summary['unmeasured_count']:
            status = 'partial'
        with (self.directory / 'blocks.csv').open('w', newline='', encoding='utf-8-sig') as stream:
            writer = csv.DictWriter(stream, fieldnames=list(blocks[0]))
            writer.writeheader()
            writer.writerows(_csv_record(row) for row in blocks)
        write_blocks_xlsx(self.directory / 'blocks.xlsx', self.plan, blocks)
        self._write_json('summary.json', dict(status=status, reason=reason, summary=summary,
                                             blocks=blocks, metadata=self.metadata))
        self._write_json('run.json', dict(status=status, reason=reason, plan=self.plan, metadata=self.metadata))
        return blocks, summary
